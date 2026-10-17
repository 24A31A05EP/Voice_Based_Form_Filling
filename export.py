"""
export.py — Excel and PDF export manager for the Voice Form Filling System.
Generates timestamped reports and saves them to the reports/ folder.
"""

import os
import logging
from datetime import datetime

import pandas as pd
from reportlab.lib             import colors
from reportlab.lib.pagesizes   import A4, landscape
from reportlab.lib.styles      import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units       import cm
from reportlab.platypus        import (
    SimpleDocTemplate, Table, TableStyle,
    Paragraph, Spacer, HRFlowable,
)
from reportlab.lib.enums       import TA_CENTER, TA_LEFT
from config import EXCEL_DIR, PDF_DIR, APP_TITLE, APP_VERSION

logger = logging.getLogger(__name__)

# Ensure output directories exist
os.makedirs(EXCEL_DIR, exist_ok=True)
os.makedirs(PDF_DIR,   exist_ok=True)

# Column headers displayed in reports
REPORT_COLUMNS = [
    ("id",         "ID"),
    ("name",       "Full Name"),
    ("age",        "Age"),
    ("gender",     "Gender"),
    ("phone",      "Mobile"),
    ("email",      "Email"),
    ("address",    "Address"),
    ("occupation", "Occupation"),
    ("reg_date",   "Reg. Date"),
    ("reg_time",   "Reg. Time"),
]


class ExportManager:
    """
    Handles exporting user records to Excel (.xlsx) and PDF formats.

    Usage
    -----
    em = ExportManager()
    path = em.export_to_excel(records)
    path = em.export_to_pdf(records)
    """

    def __init__(self):
        logger.info("ExportManager initialised.")

    # ─────────────────────────────────────────
    # Helpers
    # ─────────────────────────────────────────
    @staticmethod
    def _timestamp() -> str:
        return datetime.now().strftime("%Y%m%d_%H%M%S")

    @staticmethod
    def _records_to_df(records: list[dict]) -> pd.DataFrame:
        """Convert list of record dicts to a formatted DataFrame."""
        if not records:
            return pd.DataFrame()
        df = pd.DataFrame(records)
        # Reorder and rename columns
        keys   = [c[0] for c in REPORT_COLUMNS]
        labels = [c[1] for c in REPORT_COLUMNS]
        # Keep only known columns
        existing = [k for k in keys if k in df.columns]
        df = df[existing].copy()
        rename_map = {c[0]: c[1] for c in REPORT_COLUMNS}
        df.rename(columns=rename_map, inplace=True)
        return df

    # ─────────────────────────────────────────
    # Excel Export
    # ─────────────────────────────────────────
    def export_to_excel(self, records: list[dict], filename: str = None) -> str:
        """
        Export records to a formatted .xlsx file.

        Parameters
        ----------
        records  : list[dict]  Records from DatabaseManager.get_all_users()
        filename : str         Optional custom filename (without extension).

        Returns
        -------
        str  Absolute path to the created Excel file.
        """
        if not records:
            raise ValueError("No records to export.")

        filename = filename or f"VoiceForm_Report_{self._timestamp()}"
        path     = os.path.join(EXCEL_DIR, f"{filename}.xlsx")

        df = self._records_to_df(records)

        try:
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Registrations")

                # Style the workbook
                workbook  = writer.book
                worksheet = writer.sheets["Registrations"]

                # Auto-fit column widths
                for col in worksheet.columns:
                    max_len  = max(len(str(cell.value or "")) for cell in col)
                    col_letter = col[0].column_letter
                    worksheet.column_dimensions[col_letter].width = min(max_len + 4, 50)

                # Header style
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                header_fill  = PatternFill("solid", fgColor="1E88E5")
                header_font  = Font(bold=True, color="FFFFFF", size=11)
                border_side  = Side(style="thin", color="D0D0D0")
                cell_border  = Border(
                    left=border_side, right=border_side,
                    top=border_side,  bottom=border_side,
                )
                for cell in worksheet[1]:
                    cell.fill      = header_fill
                    cell.font      = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border    = cell_border

                # Alternate row shading
                light_fill = PatternFill("solid", fgColor="EAF3FB")
                for row_idx, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
                    fill = light_fill if row_idx % 2 == 0 else None
                    for cell in row:
                        cell.border    = cell_border
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                        if fill:
                            cell.fill = fill

                # Freeze header row
                worksheet.freeze_panes = "A2"

                # Summary sheet
                summary_ws = workbook.create_sheet("Summary")
                summary_ws["A1"] = f"{APP_TITLE} v{APP_VERSION}"
                summary_ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
                summary_ws["A3"] = f"Total Records: {len(records)}"
                summary_ws["A1"].font = Font(bold=True, size=14, color="1E88E5")

            logger.info("Excel exported: %s", path)
            return path

        except Exception as exc:
            logger.error("Excel export error: %s", exc)
            raise

    # ─────────────────────────────────────────
    # PDF Export
    # ─────────────────────────────────────────
    def export_to_pdf(self, records: list[dict], filename: str = None) -> str:
        """
        Export records to a formatted PDF report.

        Parameters
        ----------
        records  : list[dict]
        filename : str  Optional custom filename (without extension).

        Returns
        -------
        str  Absolute path to the created PDF file.
        """
        if not records:
            raise ValueError("No records to export.")

        filename = filename or f"VoiceForm_Report_{self._timestamp()}"
        path     = os.path.join(PDF_DIR, f"{filename}.pdf")

        try:
            doc = SimpleDocTemplate(
                path,
                pagesize=landscape(A4),
                leftMargin=1.5*cm,
                rightMargin=1.5*cm,
                topMargin=2*cm,
                bottomMargin=2*cm,
            )

            styles = getSampleStyleSheet()
            story  = []

            # ── Title
            title_style = ParagraphStyle(
                "Title",
                parent=styles["Title"],
                fontSize=20,
                textColor=colors.HexColor("#1E88E5"),
                alignment=TA_CENTER,
                spaceAfter=6,
            )
            story.append(Paragraph(f"📋 {APP_TITLE} — Registration Report", title_style))

            sub_style = ParagraphStyle(
                "Sub",
                parent=styles["Normal"],
                fontSize=10,
                textColor=colors.HexColor("#4A5568"),
                alignment=TA_CENTER,
                spaceAfter=4,
            )
            story.append(Paragraph(
                f"Generated on {datetime.now().strftime('%d %B %Y at %H:%M:%S')} | "
                f"Total Records: {len(records)}",
                sub_style,
            ))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#E2E8F0")))
            story.append(Spacer(1, 0.4*cm))

            # ── Build table data
            headers  = [c[1] for c in REPORT_COLUMNS]
            data     = [headers]

            for rec in records:
                row = [str(rec.get(c[0], "")) for c in REPORT_COLUMNS]
                data.append(row)

            # ── Table style
            col_count  = len(headers)
            page_width = landscape(A4)[0] - 3*cm
            col_widths = [page_width / col_count] * col_count
            # Give address more space
            addr_idx = next((i for i, c in enumerate(REPORT_COLUMNS) if c[0] == "address"), None)
            if addr_idx is not None:
                extra = 3*cm
                col_widths[addr_idx] += extra
                # Reduce from id and age cols
                col_widths[0] -= extra * 0.5
                col_widths[2] -= extra * 0.5

            table = Table(data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle([
                # Header
                ("BACKGROUND",   (0, 0), (-1, 0),  colors.HexColor("#1E88E5")),
                ("TEXTCOLOR",    (0, 0), (-1, 0),  colors.white),
                ("FONTNAME",     (0, 0), (-1, 0),  "Helvetica-Bold"),
                ("FONTSIZE",     (0, 0), (-1, 0),  9),
                ("ALIGN",        (0, 0), (-1, 0),  "CENTER"),
                ("VALIGN",       (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUND", (0, 1), (-1, -1), [
                    colors.HexColor("#F0F8FF"),
                    colors.white,
                ]),
                # Body
                ("FONTNAME",    (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE",    (0, 1), (-1, -1), 8),
                ("ALIGN",       (0, 1), (-1, -1), "LEFT"),
                # Grid
                ("GRID",        (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D8E0")),
                ("LINEBELOW",   (0, 0), (-1, 0),  1.5, colors.HexColor("#1565C0")),
                # Padding
                ("TOPPADDING",  (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING",(0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ]))

            story.append(table)
            story.append(Spacer(1, 0.5*cm))

            # ── Footer note
            footer_style = ParagraphStyle(
                "Footer",
                parent=styles["Normal"],
                fontSize=8,
                textColor=colors.HexColor("#A0AEC0"),
                alignment=TA_CENTER,
            )
            story.append(Paragraph(
                f"© {datetime.now().year} {APP_TITLE}  |  Version {APP_VERSION}  |  Confidential",
                footer_style,
            ))

            doc.build(story)
            logger.info("PDF exported: %s", path)
            return path

        except Exception as exc:
            logger.error("PDF export error: %s", exc)
            raise
